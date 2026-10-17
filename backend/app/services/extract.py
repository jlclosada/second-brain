"""Extracción de texto: PDF con texto, PDF escaneado (OCR), imágenes, Word y texto plano."""

import logging
import shutil
import subprocess
import tempfile
from dataclasses import dataclass
from pathlib import Path

from app.config import get_settings

log = logging.getLogger(__name__)

TEXT_TYPES = {".txt", ".md", ".csv", ".json", ".html", ".htm", ".eml", ".xml"}
IMAGE_TYPES = {".png", ".jpg", ".jpeg", ".tif", ".tiff", ".webp", ".bmp"}

# Por debajo de esta media de caracteres por página asumimos que el PDF es un escaneo.
MIN_CHARS_PER_PAGE = 40


@dataclass
class Extracted:
    text: str
    pages: int | None = None
    ocr_used: bool = False


def extract_text(path: Path, mime: str = "") -> Extracted:
    suffix = path.suffix.lower()
    if suffix == ".pdf" or mime == "application/pdf":
        return _pdf(path)
    if suffix in IMAGE_TYPES or mime.startswith("image/"):
        return Extracted(text=_ocr_image(path), pages=1, ocr_used=True)
    if suffix == ".docx":
        return Extracted(text=_docx(path))
    if suffix in TEXT_TYPES or mime.startswith("text/"):
        return Extracted(text=path.read_text(errors="replace"))
    return Extracted(text="")


def _pdf_text(path: Path) -> tuple[str, int]:
    from pypdf import PdfReader

    reader = PdfReader(str(path))
    pages = [page.extract_text() or "" for page in reader.pages]
    return "\n\n".join(pages).strip(), len(pages)


def _pdf(path: Path) -> Extracted:
    try:
        text, pages = _pdf_text(path)
    except Exception as exc:  # PDF dañado o cifrado
        log.warning("No se pudo leer %s: %s", path.name, exc)
        text, pages = "", 0
    if pages and len(text) / pages >= MIN_CHARS_PER_PAGE:
        return Extracted(text=text, pages=pages)

    ocr_text = _ocr_pdf(path)
    if ocr_text:
        return Extracted(text=ocr_text, pages=pages or None, ocr_used=True)
    return Extracted(text=text, pages=pages or None)


def _ocr_pdf(path: Path) -> str:
    if not shutil.which("ocrmypdf"):
        log.info("ocrmypdf no está instalado; se omite el OCR de %s", path.name)
        return ""
    with tempfile.TemporaryDirectory() as tmp:
        out = Path(tmp) / "ocr.pdf"
        result = subprocess.run(
            ["ocrmypdf", "--skip-text", "--quiet", "-l", get_settings().ocr_languages, str(path), str(out)],
            capture_output=True,
            text=True,
            timeout=600,
        )
        if result.returncode != 0 or not out.exists():
            log.warning("OCR fallido en %s: %s", path.name, result.stderr[-500:])
            return ""
        text, _ = _pdf_text(out)
        return text


def _ocr_image(path: Path) -> str:
    if not shutil.which("tesseract"):
        log.info("tesseract no está instalado; se omite el OCR de %s", path.name)
        return ""
    result = subprocess.run(
        ["tesseract", str(path), "stdout", "-l", get_settings().ocr_languages],
        capture_output=True,
        text=True,
        timeout=300,
    )
    if result.returncode != 0:
        log.warning("OCR fallido en %s: %s", path.name, result.stderr[-500:])
    return result.stdout.strip()


def _docx(path: Path) -> str:
    from docx import Document

    doc = Document(str(path))
    parts = [p.text for p in doc.paragraphs]
    for table in doc.tables:
        for row in table.rows:
            parts.append(" | ".join(cell.text for cell in row.cells))
    return "\n".join(parts).strip()
