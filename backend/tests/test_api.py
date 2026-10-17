import io
from pathlib import Path

from docx import Document as DocxDocument
from fastapi.testclient import TestClient

from app.main import app

POLIZA_COCHE_2024 = """MAPFRE SEGUROS
Póliza de seguro de automóvil a todo riesgo
Tomador: José López. Vehículo: Seat León, matrícula 1234 ABC.
Periodo de cobertura: 01/03/2024 a 28/02/2025.
Prima anual: 412,50 euros."""

POLIZA_COCHE_2023 = """LÍNEA DIRECTA
Póliza de seguro de automóvil a terceros
Vehículo: Seat León, matrícula 1234 ABC.
Periodo de cobertura: 01/03/2023 a 29/02/2024. Año 2023."""

SEGURO_HOGAR_2024 = """Seguro de hogar multirriesgo
Vivienda en Calle Mayor 5. Continente y contenido.
Vigencia: 15/06/2024. Hogar 2024."""

FACTURA_LUZ = """Factura de electricidad
Iberdrola. Periodo de facturación 01/05/2024 - 31/05/2024.
Consumo luz vivienda: 230 kWh. Total 64,20 euros."""


def test_requires_login():
    with TestClient(app, headers={"X-Requested-With": "segundo-cerebro"}) as anon:
        assert anon.get("/api/documents").status_code == 401
        assert anon.post("/api/auth/login", json={"password": "mal"}).status_code == 401


def test_upload_classifies_with_rules(client, upload, process_all):
    created = upload("scan_0032.txt", POLIZA_COCHE_2024)["created"][0]
    assert created["status"] == "pending"
    process_all()

    doc = client.get(f"/api/documents/{created['id']}").json()
    assert doc["status"] == "ready"
    assert doc["category"] == "Vehículo"
    assert doc["doc_type"] == "póliza"
    assert doc["year"] == 2024
    assert doc["doc_date"] == "2024-03-01"
    assert "Prima anual" in doc["text"]


def test_duplicate_upload_is_skipped(client, upload):
    upload("a.txt", POLIZA_COCHE_2024)
    second = upload("copia.txt", POLIZA_COCHE_2024)
    assert second["created"] == []
    assert second["duplicates"] == ["copia.txt"]


def test_search_seguro_de_coche_2024(client, upload, process_all):
    for name, body in [
        ("scan_0032.txt", POLIZA_COCHE_2024),
        ("scan_0017.txt", POLIZA_COCHE_2023),
        ("hogar.txt", SEGURO_HOGAR_2024),
        ("luz_mayo.txt", FACTURA_LUZ),
    ]:
        upload(name, body)
    process_all()

    r = client.get("/api/search", params={"q": "seguro de coche 2024"}).json()
    assert r["years"] == [2024]
    assert r["text"] == "seguro de coche"
    titles = [h["item"]["document"]["filename"] for h in r["hits"]]
    assert titles[0] == "scan_0032.txt"
    assert "scan_0017.txt" not in titles  # es de 2023
    assert "<mark>" in r["hits"][0]["snippet"]


def test_search_ignores_accents_and_falls_back_without_year(client, upload, process_all):
    upload("scan_0017.txt", POLIZA_COCHE_2023)
    process_all()
    # No hay nada de 2030: se devuelve lo más parecido sin el filtro de año.
    r = client.get("/api/search", params={"q": "poliza automovil 2030"}).json()
    assert [h["item"]["document"]["filename"] for h in r["hits"]] == ["scan_0017.txt"]


def test_search_only_year_lists_that_year(client, upload, process_all):
    upload("a.txt", POLIZA_COCHE_2024)
    upload("b.txt", POLIZA_COCHE_2023)
    process_all()
    r = client.get("/api/search", params={"q": "2023"}).json()
    assert [h["item"]["document"]["filename"] for h in r["hits"]] == ["b.txt"]


def test_edit_reindexes_card(client, upload, process_all):
    item = upload("nota.txt", "Garantía de la lavadora Bosch comprada en MediaMarkt")["created"][0]
    process_all()
    r = client.patch(f"/api/documents/{item['id']}", json={"category": "Compras", "tags": ["Electrodomésticos"]})
    assert r.status_code == 200
    assert r.json()["tags"] == ["electrodomésticos"]
    process_all()  # el reindexado va por la cola
    hits = client.get("/api/search", params={"q": "electrodomésticos"}).json()["hits"]
    assert hits and hits[0]["item"]["id"] == item["id"]
    assert client.patch(f"/api/documents/{item['id']}", json={"category": "Inventada"}).status_code == 422


def test_docx_and_file_download_and_delete(client, upload, process_all):
    buf = io.BytesIO()
    doc = DocxDocument()
    doc.add_paragraph("Contrato de alquiler de vivienda")
    doc.add_paragraph("Firmado el 10/01/2025 en Madrid.")
    doc.save(buf)
    mime = "application/vnd.openxmlformats-officedocument.wordprocessingml.document"
    item = upload("contrato.docx", buf.getvalue(), mime)["created"][0]
    process_all()

    detail = client.get(f"/api/documents/{item['id']}").json()
    assert detail["category"] == "Hogar"
    assert detail["year"] == 2025
    assert client.get(f"/api/documents/{item['id']}/file").content == buf.getvalue()

    from app.db import SessionLocal
    from app.models import Document

    with SessionLocal() as s:
        path = Path(s.get(Document, detail["id"]).path)
    assert path.exists()
    assert client.delete(f"/api/documents/{item['id']}").status_code == 204
    assert not path.exists()
    assert client.get(f"/api/documents/{item['id']}").status_code == 404


def test_stats_and_filters(client, upload, process_all):
    upload("a.txt", POLIZA_COCHE_2024)
    upload("b.txt", SEGURO_HOGAR_2024)
    process_all()
    stats = client.get("/api/stats").json()
    assert stats["total"] == 2
    assert {c["value"] for c in stats["categories"]} == {"Vehículo", "Hogar"}
    listed = client.get("/api/documents", params={"category": "Hogar"}).json()
    assert [i["document"]["filename"] for i in listed["items"]] == ["b.txt"]
    assert client.get("/api/documents", params={"status": "ready"}).json()["total"] == 2


def test_assistant_without_api_key(client):
    r = client.post("/api/assistant/chat", json={"message": "¿Cuándo vence mi seguro de coche?"})
    assert r.status_code == 503


def test_snippet_escapes_document_html(client, upload, process_all):
    upload("pagina.html", "<img src=x onerror=alert(1)> Recibo del gimnasio", "text/html")
    process_all()
    snippet = client.get("/api/search", params={"q": "gimnasio"}).json()["hits"][0]["snippet"]
    assert "<img" not in snippet
    assert "&lt;img" in snippet
    assert "<mark>gimnasio</mark>" in snippet
    item_id = client.get("/api/documents").json()["items"][0]["id"]
    r = client.get(f"/api/documents/{item_id}/file")
    assert r.headers["content-security-policy"] == "sandbox"
